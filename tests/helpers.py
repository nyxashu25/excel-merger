"""Shared helpers for building small spreadsheets in a temp folder."""
import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


def make_xlsx(path, sheets):
    """sheets: {sheet name: [row, row, ...]}; a name starting with '!' is created hidden (the '!' is stripped)."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        hidden = name.startswith("!")
        ws = wb.create_sheet(name.lstrip("!"))
        for row in rows:
            ws.append(row)
        if hidden:
            ws.sheet_state = "hidden"
    wb.save(path)
    wb.close()
    return Path(path)


def make_csv(path, rows, delimiter=",", encoding="utf-8"):
    with open(path, "w", newline="", encoding=encoding) as fh:
        csv.writer(fh, delimiter=delimiter).writerows(rows)
    return Path(path)


def read_xlsx(path):
    """Returns {sheet name: [row tuples]} with trailing empty cells kept as None."""
    wb = load_workbook(path, data_only=True)
    try:
        return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()


class TempDirCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)

    def path(self, name):
        return self.dir / name
