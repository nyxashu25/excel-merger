import contextlib
import datetime as dt
import gc
import io
import unittest
from unittest import mock

from openpyxl import load_workbook

from excel_merger import MergeError, collect_files, merge_files
from helpers import TempDirCase, make_csv, make_xlsx, read_xlsx


class StackTests(TempDirCase):
    def test_aligns_columns_by_name_and_fills_gaps(self):
        a = make_xlsx(self.path("a.xlsx"), {"Data": [["Name", "Email"], ["Ann", "ann@x.com"]]})
        b = make_xlsx(self.path("b.xlsx"), {"Data": [["Email", "Phone", "Name"], ["bob@x.com", "555", "Bob"]]})
        out = self.path("out.xlsx")
        res = merge_files([a, b], out)
        rows = read_xlsx(out)["Merged"]
        self.assertEqual(rows[0], ("Name", "Email", "Phone", "Source file", "Source sheet"))
        self.assertEqual(rows[1], ("Ann", "ann@x.com", None, "a.xlsx", "Data"))
        self.assertEqual(rows[2], ("Bob", "bob@x.com", "555", "b.xlsx", "Data"))
        self.assertEqual((res.files, res.sheets, res.rows, res.columns), (2, 2, 2, 5))

    def test_header_match_is_case_insensitive_by_default(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["Email"], ["1"]]})
        b = make_xlsx(self.path("b.xlsx"), {"S": [["  email "], ["2"]]})
        out = self.path("out.xlsx")
        merge_files([a, b], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("Email",), ("1",), ("2",)])

    def test_match_case_keeps_columns_separate(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["Email"], ["1"]]})
        b = make_xlsx(self.path("b.xlsx"), {"S": [["email"], ["2"]]})
        out = self.path("out.xlsx")
        merge_files([a, b], out, source_columns=False, match_case=True)
        self.assertEqual(read_xlsx(out)["Merged"], [("Email", "email"), ("1", None), (None, "2")])

    def test_dedupe_removes_identical_rows_across_files(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["id", "v"], [1, "x"], [2, "y"], [1, "x"]]})
        b = make_xlsx(self.path("b.xlsx"), {"S": [["id", "v"], [2, "y"], [3, "z"]]})
        out = self.path("out.xlsx")
        res = merge_files([a, b], out, dedupe=True, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("id", "v"), (1, "x"), (2, "y"), (3, "z")])
        self.assertEqual(res.duplicates_removed, 2)
        self.assertEqual(res.rows, 3)

    def test_header_row_skips_title_rows(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["Quarterly report"], [], ["id", "v"], [1, "x"]]})
        out = self.path("out.xlsx")
        merge_files([a], out, header_row=3, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("id", "v"), (1, "x")])

    def test_empty_rows_are_skipped(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["id"], [1], [None], ["  "], [2]]})
        out = self.path("out.xlsx")
        res = merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("id",), (1,), (2,)])
        self.assertEqual(res.rows, 2)

    def test_blank_and_duplicate_headers_get_unique_names(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["id", None, "id"], [1, 2, 3]]})
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("id", "Column 2", "id (2)"), (1, 2, 3)])

    def test_data_wider_than_header_is_not_lost(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["id"], [1, "extra"]]})
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("id", "Column 2"), (1, "extra")])

    def test_numbers_dates_and_bools_keep_their_types(self):
        when = dt.datetime(2024, 5, 17, 9, 30)
        a = make_xlsx(self.path("a.xlsx"), {"S": [["n", "f", "d", "b"], [7, 2.5, when, True]]})
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        row = read_xlsx(out)["Merged"][1]
        self.assertEqual(row, (7, 2.5, when, True))
        self.assertIsInstance(row[0], int)

    def test_text_starting_with_equals_never_becomes_a_formula(self):
        a = make_csv(self.path("a.csv"), [["note"], ["=1+1"], ["=HYPERLINK(\"http://evil\")"]])
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        wb = load_workbook(out)  # not data_only: formulas would show up as data_type 'f'
        ws = wb["Merged"]
        self.assertEqual(ws["A2"].value, "=1+1")
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["A3"].data_type, "s")
        wb.close()

    def test_source_columns_do_not_collide_with_existing_headers(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["Source file", "v"], ["mine", 1]]})
        out = self.path("out.xlsx")
        merge_files([a], out)
        self.assertEqual(read_xlsx(out)["Merged"][0], ("Source file", "v", "Source file (2)", "Source sheet"))

    def test_formula_cells_are_merged_as_their_cached_values(self):
        # openpyxl cannot compute formulas, so a file it writes has no cached value: the merge must not crash.
        wb_path = self.path("f.xlsx")
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["a", "b"])
        wb.active.append([1, "=A2+1"])
        wb.save(wb_path)
        out = self.path("out.xlsx")
        merge_files([wb_path], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"][1][0], 1)


class SheetSelectionTests(TempDirCase):
    def test_all_visible_sheets_by_default_and_hidden_on_request(self):
        a = make_xlsx(self.path("a.xlsx"), {"One": [["v"], [1]], "Two": [["v"], [2]], "!Secret": [["v"], [3]]})
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("v",), (1,), (2,)])
        out2 = self.path("out2.xlsx")
        merge_files([a], out2, source_columns=False, include_hidden=True)
        self.assertEqual(read_xlsx(out2)["Merged"], [("v",), (1,), (2,), (3,)])

    def test_sheet_filter_is_case_insensitive_and_warns_when_missing(self):
        a = make_xlsx(self.path("a.xlsx"), {"Data": [["v"], [1]], "Other": [["v"], [9]]})
        b = make_xlsx(self.path("b.xlsx"), {"Nothing": [["v"], [5]]})
        out = self.path("out.xlsx")
        res = merge_files([a, b], out, sheets=["data"], source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("v",), (1,)])
        self.assertTrue(any("b.xlsx" in w for w in res.warnings), res.warnings)

    def test_nothing_to_merge_is_an_error(self):
        a = make_xlsx(self.path("a.xlsx"), {"Empty": []})
        with self.assertRaises(MergeError):
            merge_files([a], self.path("out.xlsx"))
        self.assertFalse(self.path("out.xlsx").exists(), "no partial output on failure")


class CsvTests(TempDirCase):
    def test_semicolon_delimiter_is_detected(self):
        a = make_csv(self.path("a.csv"), [["id", "name"], ["1", "Ann"], ["2", "Bob"]], delimiter=";")
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("id", "name"), ("1", "Ann"), ("2", "Bob")])

    def test_explicit_delimiter(self):
        a = make_csv(self.path("a.csv"), [["id", "name"], ["1", "Ann"]], delimiter="|")
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False, delimiter="|")
        self.assertEqual(read_xlsx(out)["Merged"], [("id", "name"), ("1", "Ann")])

    def test_windows_1252_files_are_readable(self):
        a = make_csv(self.path("a.csv"), [["name"], ["Café"], ["Zoë"]], encoding="cp1252")
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("name",), ("Café",), ("Zoë",)])

    def test_utf8_bom_is_stripped_from_the_first_header(self):
        a = make_csv(self.path("a.csv"), [["id"], ["1"]], encoding="utf-8-sig")
        out = self.path("out.xlsx")
        merge_files([a], out, source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"][0], ("id",))

    def test_csv_and_xlsx_stack_together_and_csv_output_round_trips(self):
        a = make_csv(self.path("a.csv"), [["id", "v"], ["1", "x"]])
        b = make_xlsx(self.path("b.xlsx"), {"S": [["id", "v"], [2, "y"]]})
        out = self.path("out.csv")
        merge_files([a, b], out)
        text = out.read_text(encoding="utf-8-sig").splitlines()
        self.assertEqual(text[0], "id,v,Source file,Source sheet")
        self.assertEqual(text[1], "1,x,a.csv,a")
        self.assertEqual(text[2], "2,y,b.xlsx,S")

    def test_sheet_filter_does_not_exclude_csv_files(self):
        a = make_csv(self.path("a.csv"), [["v"], ["1"]])
        b = make_xlsx(self.path("b.xlsx"), {"Keep": [["v"], [2]], "Drop": [["v"], [3]]})
        out = self.path("out.xlsx")
        merge_files([a, b], out, sheets=["Keep"], source_columns=False)
        self.assertEqual(read_xlsx(out)["Merged"], [("v",), ("1",), (2,)])


class SheetsModeTests(TempDirCase):
    def test_one_output_sheet_per_input_sheet(self):
        a = make_xlsx(self.path("jan.xlsx"), {"Sales": [["v"], [1]]})
        b = make_xlsx(self.path("feb.xlsx"), {"Sales": [["v"], [2]], "Costs": [["c"], [3]]})
        out = self.path("book.xlsx")
        res = merge_files([a, b], out, mode="sheets")
        data = read_xlsx(out)
        self.assertEqual(list(data), ["jan", "feb_Sales", "feb_Costs"])
        self.assertEqual(data["jan"], [("v",), (1,)])
        self.assertEqual(data["feb_Costs"], [("c",), (3,)])
        self.assertEqual((res.sheets, res.rows), (3, 3))

    def test_sheet_names_are_sanitized_truncated_and_unique(self):
        long_name = "x" * 40
        a = make_xlsx(self.path(f"{long_name}.xlsx"), {"S": [["v"], [1]]})
        sub = self.path("sub")
        sub.mkdir()
        b = make_xlsx(sub / f"{long_name}.xlsx", {"S": [["v"], [2]]})
        out = self.path("book.xlsx")
        merge_files([a, b], out, mode="sheets")
        names = list(read_xlsx(out))
        self.assertEqual(len(names), 2)
        self.assertTrue(all(len(n) <= 31 for n in names), names)
        self.assertEqual(len({n.casefold() for n in names}), 2)

    def test_invalid_combinations_are_rejected(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"], [1]]})
        with self.assertRaises(MergeError):
            merge_files([a], self.path("o.csv"), mode="sheets")
        with self.assertRaises(MergeError):
            merge_files([a], self.path("o.xlsx"), mode="sheets", dedupe=True)


class InputAndOutputSafetyTests(TempDirCase):
    def test_folder_input_is_naturally_sorted_and_skips_lock_and_hidden_files(self):
        folder = self.path("in")
        folder.mkdir()
        for n in (10, 2, 1):
            make_xlsx(folder / f"report{n}.xlsx", {"S": [["v"], [n]]})
        make_xlsx(folder / "~$report1.xlsx", {"S": [["v"], [99]]})
        make_xlsx(folder / ".hidden.xlsx", {"S": [["v"], [98]]})
        (folder / "notes.txt").write_text("ignore me")
        names = [p.name for p in collect_files([folder])]
        self.assertEqual(names, ["report1.xlsx", "report2.xlsx", "report10.xlsx"])

    def test_recursive_flag(self):
        folder = self.path("in")
        (folder / "deep").mkdir(parents=True)
        make_xlsx(folder / "a.xlsx", {"S": [["v"], [1]]})
        make_xlsx(folder / "deep" / "b.xlsx", {"S": [["v"], [2]]})
        self.assertEqual([p.name for p in collect_files([folder])], ["a.xlsx"])
        self.assertEqual(sorted(p.name for p in collect_files([folder], recursive=True)), ["a.xlsx", "b.xlsx"])

    def test_glob_patterns(self):
        make_xlsx(self.path("a1.xlsx"), {"S": [["v"], [1]]})
        make_xlsx(self.path("a2.xlsx"), {"S": [["v"], [2]]})
        make_xlsx(self.path("b1.xlsx"), {"S": [["v"], [3]]})
        names = [p.name for p in collect_files([str(self.path("a*.xlsx"))])]
        self.assertEqual(names, ["a1.xlsx", "a2.xlsx"])
        with self.assertRaises(MergeError):
            collect_files([str(self.path("zzz*.xlsx"))])

    def test_rerunning_into_the_input_folder_does_not_merge_the_previous_output(self):
        folder = self.path("in")
        folder.mkdir()
        make_xlsx(folder / "a.xlsx", {"S": [["v"], [1]]})
        out = folder / "merged.xlsx"
        merge_files([folder], out, source_columns=False)
        res = merge_files([folder], out, source_columns=False, overwrite=True)
        self.assertEqual(res.files, 1)
        self.assertEqual(read_xlsx(out)["Merged"], [("v",), (1,)])

    def test_refuses_to_overwrite_without_flag_and_leaves_the_file_alone(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"], [1]]})
        out = make_xlsx(self.path("out.xlsx"), {"Precious": [["keep"], ["me"]]})
        with self.assertRaises(MergeError):
            merge_files([a], out)
        self.assertEqual(read_xlsx(out), {"Precious": [("keep",), ("me",)]})
        merge_files([a], out, overwrite=True)
        self.assertIn("Merged", read_xlsx(out))

    def test_output_that_is_also_an_input_is_rejected(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"], [1]]})
        with self.assertRaises(MergeError):
            merge_files([a], a, overwrite=True)

    def test_no_temp_files_are_left_behind(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"], [1]]})
        merge_files([a], self.path("out.xlsx"))
        with self.assertRaises(MergeError):
            merge_files([a], self.path("out2.xlsx"), header_row=50)
        leftovers = [p.name for p in self.dir.iterdir() if p.name.endswith(".tmp.xlsx") or ".tmp" in p.suffixes]
        self.assertEqual(leftovers, [])

    def test_failure_while_writing_leaves_no_output_and_no_temp_file(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"]] + [[i] for i in range(10)]})
        out = self.path("out.xlsx")
        with mock.patch("excel_merger.core.MAX_EXCEL_ROWS", 5):
            with self.assertRaisesRegex(MergeError, "row limit"):
                merge_files([a], out)
        self.assertFalse(out.exists())
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["a.xlsx"])

    def test_row_limit_with_dedupe_fails_cleanly_without_writer_noise(self):
        # Unique rows, so dedupe cannot shrink them: the limit is only discovered mid-write.
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"]] + [[i] for i in range(10)]})
        out = self.path("out.xlsx")
        noise = io.StringIO()
        with contextlib.redirect_stderr(noise), mock.patch("excel_merger.core.MAX_EXCEL_ROWS", 5):
            with self.assertRaisesRegex(MergeError, "row limit"):
                merge_files([a], out, dedupe=True)
            gc.collect()
        self.assertEqual(noise.getvalue(), "", "abandoned openpyxl writers must not spill tracebacks")
        self.assertFalse(out.exists())
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["a.xlsx"])

    def test_row_limit_error_in_sheets_mode(self):
        a = make_xlsx(self.path("a.xlsx"), {"S": [["v"]] + [[i] for i in range(10)]})
        with mock.patch("excel_merger.core.MAX_EXCEL_ROWS", 5):
            with self.assertRaisesRegex(MergeError, "more rows"):
                merge_files([a], self.path("out.xlsx"), mode="sheets")
        self.assertEqual(sorted(p.name for p in self.dir.iterdir()), ["a.xlsx"])

    def test_helpful_errors(self):
        legacy = self.path("old.xls")
        legacy.write_bytes(b"\xd0\xcf\x11\xe0")
        with self.assertRaisesRegex(MergeError, r"\.xls"):
            collect_files([legacy])
        with self.assertRaisesRegex(MergeError, "Not found"):
            collect_files([self.path("missing.xlsx")])
        bad = self.path("broken.xlsx")
        bad.write_text("this is not a zip")
        with self.assertRaisesRegex(MergeError, "broken.xlsx"):
            merge_files([bad], self.path("out.xlsx"))
        with self.assertRaises(MergeError):
            merge_files([make_xlsx(self.path("a.xlsx"), {"S": [["v"], [1]]})], self.path("out.docx"))


if __name__ == "__main__":
    unittest.main()
