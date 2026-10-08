"""Core merge logic for excel-merger.

Files are read in two streaming passes. Pass one scans every sheet to learn its header and size so the
output columns can be aligned across files. Pass two streams the rows into a write-only workbook (or CSV),
so memory stays flat no matter how large the inputs are (apart from the optional de-duplication set).
"""
from __future__ import annotations

import codecs
import csv
import glob
import os
import re
import tempfile
from contextlib import closing, contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, Iterator, List, Optional, Sequence, Tuple, Union
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm")
CSV_SUFFIXES = (".csv",)
SUPPORTED_SUFFIXES = EXCEL_SUFFIXES + CSV_SUFFIXES
MAX_EXCEL_ROWS = 1_048_576
SHEET_NAME_MAX = 31
SOURCE_FILE_HEADER = "Source file"
SOURCE_SHEET_HEADER = "Source sheet"

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
PathLike = Union[str, "os.PathLike[str]"]


class MergeError(Exception):
    """A problem the user can fix: bad input, unreadable file, unsafe output path."""


@dataclass
class SheetInfo:
    path: Path
    sheet: str
    header: List[str]
    rows: int

    @property
    def width(self) -> int:
        return len(self.header)


@dataclass
class MergeResult:
    output: Path
    mode: str
    files: int = 0
    sheets: int = 0
    rows: int = 0
    columns: int = 0
    duplicates_removed: int = 0
    warnings: List[str] = field(default_factory=list)


# ---------------------------------------------------------------------------------------- inputs
def _natural_key(path: Path) -> List[Any]:
    parts = re.split(r"(\d+)", str(path).lower())
    return [int(p) if i % 2 else p for i, p in enumerate(parts)]


def _is_candidate(path: Path) -> bool:
    name = path.name
    return (
        path.suffix.lower() in SUPPORTED_SUFFIXES
        and not name.startswith("~$")  # Excel lock files
        and not name.startswith(".")
    )


def collect_files(
    inputs: Sequence[PathLike], recursive: bool = False, exclude: Sequence[PathLike] = ()
) -> List[Path]:
    """Expands files, directories and glob patterns into an ordered, de-duplicated list of supported files."""
    excluded = {Path(p).resolve() for p in exclude}
    found: List[Path] = []
    seen = set()

    def add(p: Path) -> None:
        resolved = p.resolve()
        if resolved in seen or resolved in excluded:
            return
        seen.add(resolved)
        found.append(p)

    for raw in inputs:
        text = os.fspath(raw)
        p = Path(text)
        if p.is_dir():
            children = p.rglob("*") if recursive else p.iterdir()
            for child in sorted(children, key=_natural_key):
                if child.is_file() and _is_candidate(child):
                    add(child)
        elif p.is_file():
            if p.name.startswith("~$"):
                continue
            suffix = p.suffix.lower()
            if suffix == ".xls":
                raise MergeError(
                    f"{p.name}: legacy .xls files are not supported. "
                    "Open it in Excel and save it as .xlsx first."
                )
            if suffix not in SUPPORTED_SUFFIXES:
                raise MergeError(f"{p.name}: unsupported file type (use .xlsx, .xlsm or .csv).")
            add(p)
        elif any(ch in text for ch in "*?["):
            hits = [Path(m) for m in sorted(glob.glob(text, recursive=True)) if Path(m).is_file()]
            hits = sorted((h for h in hits if _is_candidate(h)), key=_natural_key)
            if not hits:
                raise MergeError(f"No supported files match the pattern: {text}")
            for h in hits:
                add(h)
        else:
            raise MergeError(f"Not found: {text}")

    if not found:
        raise MergeError("No .xlsx, .xlsm or .csv files found in the given inputs.")
    return found


# ---------------------------------------------------------------------------------------- reading
def _detect_csv_encoding(path: Path) -> str:
    decoder = codecs.getincrementaldecoder("utf-8-sig")()
    try:
        with open(path, "rb") as fh:
            for chunk in iter(lambda: fh.read(1 << 20), b""):
                decoder.decode(chunk)
            decoder.decode(b"", final=True)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def _sniff_delimiter(path: Path, encoding: str) -> str:
    with open(path, newline="", encoding=encoding, errors="replace") as fh:
        sample = fh.read(8192)
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _csv_rows(path: Path, delimiter: Optional[str]) -> Iterator[List[Any]]:
    encoding = _detect_csv_encoding(path)
    delim = delimiter or _sniff_delimiter(path, encoding)
    with open(path, newline="", encoding=encoding) as fh:
        for row in csv.reader(fh, delimiter=delim):
            yield [None if cell == "" else cell for cell in row]


@contextmanager
def _open_workbook(path: Path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError) as exc:
        raise MergeError(f"Could not read {path.name}: {exc}") from exc
    try:
        yield wb
    finally:
        wb.close()


def _wanted_sheet(name: str, wanted: Optional[Sequence[str]]) -> bool:
    return not wanted or name.casefold() in {w.casefold() for w in wanted}


def _file_sheets(
    path: Path,
    wanted: Optional[Sequence[str]],
    include_hidden: bool,
    delimiter: Optional[str],
    warnings: Optional[List[str]] = None,
) -> Iterator[Tuple[str, Iterator[Sequence[Any]]]]:
    """Yields (sheet name, row iterator) for each sheet to merge. The workbook stays open while iterating."""
    if path.suffix.lower() in CSV_SUFFIXES:
        yield path.stem, _csv_rows(path, delimiter)  # --sheet only filters Excel workbooks
        return
    with _open_workbook(path) as wb:
        matched = 0
        for ws in wb.worksheets:
            if not _wanted_sheet(ws.title, wanted):
                continue
            if ws.sheet_state != "visible" and not include_hidden:
                continue
            matched += 1
            yield ws.title, ws.iter_rows(values_only=True)
        if wanted and not matched and warnings is not None:
            warnings.append(f"{path.name}: none of the requested sheets ({', '.join(wanted)}) were found.")


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _row_is_empty(row: Sequence[Any]) -> bool:
    return all(_blank(v) for v in row)


def _used_width(row: Sequence[Any]) -> int:
    for i in range(len(row) - 1, -1, -1):
        if not _blank(row[i]):
            return i + 1
    return 0


def _header_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _make_header(raw: Sequence[Any], width: int, match_case: bool) -> List[str]:
    names: List[str] = []
    seen = set()
    for i in range(width):
        name = _header_text(raw[i]) if i < len(raw) else ""
        name = name or f"Column {i + 1}"
        base, n = name, 2
        while _key(name, match_case) in seen:
            name = f"{base} ({n})"
            n += 1
        seen.add(_key(name, match_case))
        names.append(name)
    return names


def _key(name: str, match_case: bool) -> str:
    return name if match_case else name.casefold()


def _scan(path: Path, sheet: str, rows: Iterator[Sequence[Any]], header_row: int, match_case: bool) -> Optional[SheetInfo]:
    header_raw: Optional[Sequence[Any]] = None
    width = 0
    count = 0
    for idx, row in enumerate(rows, start=1):
        if idx < header_row:
            continue
        if idx == header_row:
            header_raw = row
            width = _used_width(row)
            continue
        if _row_is_empty(row):
            continue
        count += 1
        width = max(width, _used_width(row))
    if header_raw is None or width == 0:
        return None  # sheet is empty or shorter than the header row
    return SheetInfo(path, sheet, _make_header(header_raw, width, match_case), count)


def _data_rows(rows: Iterator[Sequence[Any]], header_row: int, width: int) -> Iterator[List[Any]]:
    for idx, row in enumerate(rows, start=1):
        if idx <= header_row or _row_is_empty(row):
            continue
        out = list(row[:width])
        if len(out) < width:
            out.extend([None] * (width - len(out)))
        yield out


# ---------------------------------------------------------------------------------------- writing
def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _out_value(ws: Any, value: Any) -> Any:
    """Values for an Excel sheet. Text beginning with '=' stays text so merged data can never become a live formula."""
    value = _clean(value)
    if isinstance(value, str) and value.startswith("="):
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"
        return cell
    return value


def _header_cells(ws: Any, names: Sequence[str]) -> List[Any]:
    cells = []
    for name in names:
        cell = WriteOnlyCell(ws, value=_clean(name))
        cell.font = Font(bold=True)
        cells.append(cell)
    return cells


def _set_widths(ws: Any, names: Sequence[str]) -> None:
    for i, name in enumerate(names, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(name) + 4, 12), 50)


def _sheet_title(name: str, used: set) -> str:
    title = _BAD_SHEET_CHARS.sub("_", name).strip("'").strip() or "Sheet"
    title = title[:SHEET_NAME_MAX]
    base, n = title, 2
    while title.casefold() in used:
        suffix = f"_{n}"
        title = base[: SHEET_NAME_MAX - len(suffix)] + suffix
        n += 1
    used.add(title.casefold())
    return title


def _unique_name(name: str, taken: set) -> str:
    candidate, n = name, 2
    while candidate.casefold() in taken:
        candidate = f"{name} ({n})"
        n += 1
    return candidate


def _atomic_target(output: Path) -> Path:
    fd, tmp = tempfile.mkstemp(prefix=output.stem + ".", suffix=".tmp" + output.suffix, dir=str(output.parent))
    os.close(fd)
    return Path(tmp)


# ---------------------------------------------------------------------------------------- merge
def merge_files(
    inputs: Sequence[PathLike],
    output: PathLike,
    *,
    mode: str = "stack",
    sheets: Optional[Sequence[str]] = None,
    header_row: int = 1,
    recursive: bool = False,
    source_columns: bool = True,
    dedupe: bool = False,
    match_case: bool = False,
    include_hidden: bool = False,
    delimiter: Optional[str] = None,
    overwrite: bool = False,
    on_sheet: Optional[Callable[[SheetInfo], None]] = None,
) -> MergeResult:
    """Merges spreadsheets into one file.

    mode="stack"  : stacks rows from every sheet into a single sheet, aligning columns by header name.
    mode="sheets" : copies every source sheet into its own sheet of the output workbook.
    """
    if mode not in ("stack", "sheets"):
        raise MergeError(f"Unknown mode: {mode!r} (use 'stack' or 'sheets').")
    if header_row < 1:
        raise MergeError("--header-row must be 1 or greater.")
    out_path = Path(output)
    out_suffix = out_path.suffix.lower()
    if out_suffix not in (".xlsx", ".csv"):
        raise MergeError("The output file must end in .xlsx or .csv.")
    if mode == "sheets" and out_suffix == ".csv":
        raise MergeError("--mode sheets needs an .xlsx output (CSV has no sheets).")
    if mode == "sheets" and dedupe:
        raise MergeError("--dedupe only works with --mode stack.")
    if out_path.exists() and not overwrite:
        raise MergeError(f"{out_path} already exists. Choose another name or pass --overwrite.")
    if not out_path.parent.exists():
        raise MergeError(f"The output folder does not exist: {out_path.parent}")

    files = collect_files(inputs, recursive=recursive, exclude=[out_path])
    if out_path.resolve() in {f.resolve() for f in files}:
        raise MergeError("The output file is also one of the inputs.")

    result = MergeResult(output=out_path, mode=mode)

    # Pass 1: scan headers and sizes.
    infos: List[SheetInfo] = []
    for path in files:
        with closing(_file_sheets(path, sheets, include_hidden, delimiter, result.warnings)) as stream:
            for sheet, rows in stream:
                info = _scan(path, sheet, rows, header_row, match_case)
                if info is None:
                    result.warnings.append(f"{path.name} [{sheet}]: empty, skipped.")
                    continue
                infos.append(info)
    if not infos:
        raise MergeError("Nothing to merge: no sheet with a header row and data was found.")
    result.files = len({i.path for i in infos})
    result.sheets = len(infos)

    tmp = _atomic_target(out_path)
    try:
        if mode == "stack":
            _write_stack(infos, tmp, out_suffix, result, sheets, header_row, source_columns, dedupe, match_case,
                         include_hidden, delimiter, on_sheet)
        else:
            _write_sheets(infos, tmp, result, sheets, header_row, include_hidden, delimiter, on_sheet)
        os.replace(tmp, out_path)
    except BaseException:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise
    return result


def _rows_for(info: SheetInfo, sheets, header_row, include_hidden, delimiter) -> Iterator[List[Any]]:
    """Re-opens one file and streams the data rows of a single sheet (pass 2)."""
    with closing(_file_sheets(info.path, sheets, include_hidden, delimiter)) as stream:
        for sheet, rows in stream:
            if sheet == info.sheet:
                yield from _data_rows(rows, header_row, info.width)
                return


def _write_stack(infos, tmp, suffix, result, sheets, header_row, source_columns, dedupe, match_case,
                 include_hidden, delimiter, on_sheet) -> None:
    columns: List[str] = []
    index: Dict[str, int] = {}
    for info in infos:
        for name in info.header:
            k = _key(name, match_case)
            if k not in index:
                index[k] = len(columns)
                columns.append(name)
    data_width = len(columns)
    out_columns = list(columns)
    if source_columns:
        taken = {c.casefold() for c in columns}
        file_col = _unique_name(SOURCE_FILE_HEADER, taken)
        taken.add(file_col.casefold())
        out_columns += [file_col, _unique_name(SOURCE_SHEET_HEADER, taken)]
    result.columns = len(out_columns)

    limit_msg = (
        f"The merged data exceeds Excel's {MAX_EXCEL_ROWS:,}-row limit. "
        "Write a .csv output instead, or merge fewer files."
    )
    if suffix == ".xlsx" and not dedupe and sum(i.rows for i in infos) + 1 > MAX_EXCEL_ROWS:
        raise MergeError(limit_msg)  # row counts are exact without dedupe, so fail before doing any work

    wb = None
    csv_fh = None
    overflow = False
    try:
        if suffix == ".xlsx":
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Merged")
            ws.freeze_panes = "A2"
            _set_widths(ws, out_columns)
            ws.append(_header_cells(ws, out_columns))
            emit = lambda row: ws.append([_out_value(ws, v) for v in row])  # noqa: E731
        else:
            csv_fh = open(tmp, "w", newline="", encoding="utf-8-sig")
            writer = csv.writer(csv_fh)
            writer.writerow(out_columns)
            emit = writer.writerow  # noqa: E731

        seen = set()
        for info in infos:
            if overflow:
                break
            positions = [index[_key(n, match_case)] for n in info.header]
            kept = 0
            for src in _rows_for(info, sheets, header_row, include_hidden, delimiter):
                row: List[Any] = [None] * data_width
                for pos, value in zip(positions, src):
                    row[pos] = None if _blank(value) else value
                if dedupe:
                    sig = tuple(row)
                    if sig in seen:
                        result.duplicates_removed += 1
                        continue
                    seen.add(sig)
                if source_columns:
                    row += [info.path.name, info.sheet]
                if suffix == ".xlsx" and result.rows + 2 > MAX_EXCEL_ROWS:
                    # With --dedupe the final count is unknown up front. Stop, finish the file so the
                    # writer shuts down cleanly, then report the error below.
                    overflow = True
                    break
                emit(row)
                result.rows += 1
                kept += 1
            if on_sheet:
                on_sheet(SheetInfo(info.path, info.sheet, info.header, kept))
        if wb is not None:
            wb.save(tmp)
        if overflow:
            raise MergeError(limit_msg)
    finally:
        if csv_fh is not None:
            csv_fh.close()
        if wb is not None:
            wb.close()


def _write_sheets(infos, tmp, result, sheets, header_row, include_hidden, delimiter, on_sheet) -> None:
    for info in infos:  # row counts are exact, so fail before writing anything
        if info.rows + 1 > MAX_EXCEL_ROWS:
            raise MergeError(
                f"{info.path.name} [{info.sheet}] has more rows than one Excel sheet can hold "
                f"({MAX_EXCEL_ROWS:,})."
            )
    wb = Workbook(write_only=True)
    try:
        multi = {}
        for info in infos:
            multi[info.path] = multi.get(info.path, 0) + 1
        used: set = set()
        widest = 0
        for info in infos:
            label = info.path.stem if multi[info.path] == 1 else f"{info.path.stem}_{info.sheet}"
            ws = wb.create_sheet(_sheet_title(label, used))
            ws.freeze_panes = "A2"
            _set_widths(ws, info.header)
            ws.append(_header_cells(ws, info.header))
            for row in _rows_for(info, sheets, header_row, include_hidden, delimiter):
                ws.append([_out_value(ws, v) for v in row])
                result.rows += 1
            widest = max(widest, info.width)
            if on_sheet:
                on_sheet(info)
        result.columns = widest
        wb.save(tmp)
    finally:
        wb.close()
